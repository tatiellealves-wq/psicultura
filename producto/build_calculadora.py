# -*- coding: utf-8 -*-
"""
Genera la 'Calculadora de Retorno (ROI) - Tilapia Rentable' (bono del e-book).
El usuario solo edita las celdas AMARILLAS; el resto se calcula solo.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

# ---- paleta ----
TEAL   = "0E7C86"; TEAL_D = "0A5A62"; INK = "1A2B32"
YELLOW = "FFF2B2"; YELLOW_D="D99A22"
AQUA   = "EAF4F5"; GREEN  = "E7F3EC"; GREEN_D="2E9E5B"; WHITE="FFFFFF"
GREY   = "5B6B72"

F = "Arial"
def font(sz=10, b=False, color=INK, it=False): return Font(name=F, size=sz, bold=b, color=color, italic=it)
def fill(c): return PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="CFE0E2")
med  = Side(style="medium", color=TEAL_D)
def border(**kw): return Border(**kw)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
left   = Alignment(horizontal="left", vertical="center", wrap_text=True)
right  = Alignment(horizontal="right", vertical="center")

wb = openpyxl.Workbook()

# ============================================================ HOJA 1: CALCULADORA
ws = wb.active
ws.title = "Calculadora"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 46
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 16
ws.column_dimensions["E"].width = 40
ws.column_dimensions["F"].width = 2

def title_row(r, text, sub=None):
    ws.merge_cells(f"B{r}:E{r}")
    c = ws[f"B{r}"]; c.value = text
    c.font = font(16, True, WHITE); c.fill = fill(TEAL_D); c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 30
    if sub:
        ws.merge_cells(f"B{r+1}:E{r+1}")
        c2 = ws[f"B{r+1}"]; c2.value = sub
        c2.font = font(9.5, False, WHITE, it=True); c2.fill = fill(TEAL)
        c2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r+1].height = 18

def section(r, text):
    ws.merge_cells(f"B{r}:E{r}")
    c = ws[f"B{r}"]; c.value = text
    c.font = font(11, True, WHITE); c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r].height = 22

def input_row(r, label, value, unit, note=None, money=False, pct=False, integer=False):
    ws[f"B{r}"] = label; ws[f"B{r}"].font = font(10); ws[f"B{r}"].alignment = left
    cell = ws[f"C{r}"]; cell.value = value
    cell.font = font(10, True, INK); cell.fill = fill(YELLOW)
    cell.alignment = center
    cell.border = Border(left=Side("thin",color=YELLOW_D),right=Side("thin",color=YELLOW_D),
                         top=Side("thin",color=YELLOW_D),bottom=Side("thin",color=YELLOW_D))
    if money: cell.number_format = '#,##0'
    elif pct: cell.number_format = '0%'
    elif integer: cell.number_format = '#,##0'
    ws[f"D{r}"] = unit; ws[f"D{r}"].font = font(9, color=GREY); ws[f"D{r}"].alignment = center
    if note:
        ws[f"E{r}"] = note; ws[f"E{r}"].font = font(8.5, color=GREY, it=True); ws[f"E{r}"].alignment = left
    for col in "BCDE": ws[f"{col}{r}"].border = box
    ws.row_dimensions[r].height = 20

def calc_row(r, label, formula, fmt='#,##0', bold=False, note=None, highlight=False):
    ws[f"B{r}"] = label
    ws[f"B{r}"].font = font(10, bold, TEAL_D if highlight else INK); ws[f"B{r}"].alignment = left
    cell = ws[f"C{r}"]; cell.value = formula
    cell.font = font(10, True if (bold or highlight) else False, INK)
    cell.alignment = center; cell.number_format = fmt
    if highlight: cell.fill = fill(GREEN)
    ws[f"D{r}"].font = font(9, color=GREY)
    if note:
        ws[f"E{r}"] = note; ws[f"E{r}"].font = font(8.5, color=GREY, it=True); ws[f"E{r}"].alignment = left
    for col in "BCDE":
        ws[f"{col}{r}"].border = box
        if highlight and col in "DE": ws[f"{col}{r}"].fill = fill(GREEN)
    ws.row_dimensions[r].height = 20

r = 2
title_row(r, "Calculadora de Retorno (ROI) — Tilapia Rentable",
          "Bono del e-book. Edita SOLO las celdas amarillas con los precios de tu ciudad y moneda.")
r = 5

# leyenda
ws.merge_cells(f"B{r}:E{r}")
c = ws[f"B{r}"]
c.value = "🟡 Celdas AMARILLAS = tú las editas   ·   ⬜ Celdas blancas/verdes = se calculan solas (no las toques)"
c.font = font(9, True, YELLOW_D); c.fill = fill("FBF3E2")
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
c.border = Border(left=Side("thin",color=YELLOW_D),right=Side("thin",color=YELLOW_D),
                  top=Side("thin",color=YELLOW_D),bottom=Side("thin",color=YELLOW_D))
ws.row_dimensions[r].height = 22
r += 2

# --- moneda ---
section(r, "1 · Tu moneda"); r += 1
input_row(r, "Nombre de tu moneda (ej: MXN, COP, PEN, USD)", "USD", "texto",
          "Solo referencia visual. Usa la misma moneda en toda la planilla."); r += 2

# --- Inversión fija ---
inv_fija_start = r
section(r, "2 · Inversión fija (una sola vez, dura varios ciclos)"); r += 1
input_row(r, "Tanque / estanque + estructura", 1600, "moneda", money=True); f1=r; r+=1
input_row(r, "Sistema de aireación / bomba", 650, "moneda", money=True); f2=r; r+=1
input_row(r, "Equipo de manejo, red, báscula, cosecha", 400, "moneda", money=True); f3=r; r+=1
input_row(r, "Otros equipos fijos", 0, "moneda", money=True); f4=r; r+=1
calc_row(r, "➤ Subtotal inversión fija", f"=SUM(C{f1}:C{f4})", bold=True); SUB_FIJA=r; r+=2

# --- Producción ---
section(r, "3 · Tu producción (por ciclo)"); r += 1
input_row(r, "Alevines sembrados", 3000, "peces", integer=True); p_alev=r; r+=1
input_row(r, "Mortalidad esperada", 0.10, "%", "Para principiante usa 10%.", pct=True); p_mort=r; r+=1
input_row(r, "Peso promedio a la cosecha", 0.55, "kg/pez", "Talla comercial: 0,4–0,6 kg."); p_peso=r; r+=1
calc_row(r, "➤ Peces cosechados", f"=ROUND(C{p_alev}*(1-C{p_mort}),0)", '#,##0', note="Sobreviven al ciclo."); p_cos=r; r+=1
calc_row(r, "➤ Kilos producidos por ciclo", f"=ROUND(C{p_cos}*C{p_peso},0)", '#,##0', bold=True, note="Tu cosecha total en kg."); p_kg=r; r+=2

# --- Costos de operación ---
section(r, "4 · Costo de operación (cada ciclo)"); r += 1
input_row(r, "Costo de los alevines", 330, "moneda", money=True); o1=r; r+=1
input_row(r, "Conversión alimenticia (FCA)", 1.6, "kg alim./kg pez", "Meta: ≤1,6. Es el número clave."); o_fca=r; r+=1
input_row(r, "Precio del alimento balanceado", 0.75, "moneda/kg", "Precio por kg de alimento."); o_palim=r; r+=1
calc_row(r, "➤ Costo de alimento del ciclo", f"=ROUND(C{p_kg}*C{o_fca}*C{o_palim},0)", '#,##0', bold=True,
         note="Kilos × FCA × precio. Suele ser el 40–70% del costo."); o_alim=r; r+=1
input_row(r, "Energía (todo el ciclo)", 500, "moneda", money=True); o2=r; r+=1
input_row(r, "Mano de obra (todo el ciclo)", 400, "moneda", money=True); o3=r; r+=1
input_row(r, "Otros (agua, sanidad, transporte)", 200, "moneda", money=True); o4=r; r+=1
input_row(r, "Imprevistos", 0.10, "%", "Sobre el resto de la operación.", pct=True); o_imp=r; r+=1
calc_row(r, "➤ Subtotal costo de operación",
         f"=ROUND((C{o1}+C{o_alim}+C{o2}+C{o3}+C{o4})*(1+C{o_imp}),0)", bold=True); SUB_OPER=r; r+=2

# --- Venta ---
section(r, "5 · Tu venta"); r += 1
input_row(r, "Precio de venta por kg", 3.0, "moneda/kg", "Lo que te paga el comprador."); v_precio=r; r+=1
calc_row(r, "➤ Ingreso por ciclo", f"=ROUND(C{p_kg}*C{v_precio},0)", '#,##0', bold=True); INGRESO=r; r+=2

# ============================ RESULTADOS ============================
ws.merge_cells(f"B{r}:E{r}")
c = ws[f"B{r}"]; c.value = "📊  TUS RESULTADOS"
c.font = font(13, True, WHITE); c.fill = fill(GREEN_D)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[r].height = 26; r += 1

calc_row(r, "Inversión total para ARRANCAR (fija + 1er ciclo de operación)",
         f"=C{SUB_FIJA}+C{SUB_OPER}", '#,##0', highlight=True,
         note="El dinero que necesitas ANTES de sembrar."); R_ARRANQUE=r; r+=1
calc_row(r, "Costo real por kilo producido",
         f"=ROUND(C{SUB_OPER}/C{p_kg},2)", '#,##0.00', highlight=True,
         note="Si tu precio de venta es MENOR a esto, PIERDES."); R_COSTOKG=r; r+=1
calc_row(r, "Ganancia neta por ciclo (desde el 2° ciclo)",
         f"=C{INGRESO}-C{SUB_OPER}", '#,##0', highlight=True,
         note="Ingreso menos operación (la inversión fija ya está pagada)."); R_GAN=r; r+=1
calc_row(r, "Margen sobre ventas",
         f"=IFERROR((C{INGRESO}-C{SUB_OPER})/C{INGRESO},0)", '0%', highlight=True,
         note="Qué % de cada venta te queda."); R_MARGEN=r; r+=1
calc_row(r, "Retorno del 1er ciclo (ROI incluyendo inversión fija)",
         f"=IFERROR((C{INGRESO}-C{SUB_OPER}-C{SUB_FIJA})/C{R_ARRANQUE},0)", '0%', highlight=True,
         note="Negativo normal el 1er ciclo: aún pagas la estructura."); R_ROI1=r; r+=1
calc_row(r, "Ciclos para recuperar la inversión fija",
         f"=IFERROR(ROUNDUP(C{SUB_FIJA}/(C{INGRESO}-C{SUB_OPER}),1),0)", '#,##0.0', highlight=True,
         note="Cada ciclo ≈ 6 meses."); R_RECUP=r; r+=1
calc_row(r, "⚠ Capital mínimo que debes tener asegurado",
         f"=C{SUB_FIJA}+C{SUB_OPER}", '#,##0', highlight=True,
         note="Nunca siembres sin tener cubierto este monto."); r+=2

# nota de fórmula clave con formato condicional simple (texto)
ws.merge_cells(f"B{r}:E{r}")
c = ws[f"B{r}"]
c.value = ("Regla de oro: si el «Costo real por kilo» es mayor o igual que tu «Precio de venta por kg», "
           "el proyecto pierde dinero. Baja tu FCA, negocia mejor el alimento o sube tu precio de venta.")
c.font = font(9, True, TEAL_D); c.fill = fill(AQUA); c.alignment = left
c.border = box
ws.row_dimensions[r].height = 34; r+=1
ws.merge_cells(f"B{r}:E{r}")
c = ws[f"B{r}"]
c.value = "Cifras precargadas = Escenario B del e-book (ilustrativas). Reemplázalas por las de tu ciudad."
c.font = font(8.5, color=GREY, it=True); c.alignment = left

# comentario en celda clave
ws[f"C{o_fca}"].comment = Comment(
    "FCA = kilos de alimento para producir 1 kg de pez. Es EL número del negocio. "
    "Bien manejado ronda 1,4–1,6. Arriba de 1,8 tu margen desaparece.", "Tilapia Rentable")

# ============================================================ HOJA 2: INSTRUCCIONES
ws2 = wb.create_sheet("Cómo usar")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 2
ws2.column_dimensions["B"].width = 100

def h(r, t, sz=14, color=TEAL_D, b=True, fillc=None):
    ws2.merge_cells(f"B{r}:B{r}")
    c = ws2[f"B{r}"]; c.value=t; c.font=font(sz,b,color); c.alignment=Alignment(wrap_text=True, vertical="center")
    if fillc: c.fill=fill(fillc)
    return r+1

def p(r, t, sz=10, b=False, color=INK):
    c = ws2[f"B{r}"]; c.value=t; c.font=font(sz,b,color)
    c.alignment=Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[r].height = 15*(1+len(t)//95)
    return r+1

rr = 2
rr = h(rr, "Cómo usar tu Calculadora de ROI", 18); rr+=1
rr = p(rr, "Esta planilla es el bono de la guía «Tilapia Rentable». Sirve para saber, con TUS números, "
       "si tu proyecto de tilapia es rentable y cuánto capital necesitas antes de empezar.", 11); rr+=1
rr = h(rr, "Paso 1 — Haz tu copia", 12)
rr = p(rr, "• En Excel: guarda el archivo con otro nombre antes de editar.")
rr = p(rr, "• En Google Sheets: Archivo → Hacer una copia. Así trabajas sobre TU copia y conservas la original.")
rr+=1
rr = h(rr, "Paso 2 — Edita solo las celdas amarillas", 12)
rr = p(rr, "En la hoja «Calculadora», las celdas amarillas son las únicas que debes cambiar. "
       "Reemplaza cada valor por el precio real de tu ciudad y en tu moneda (alevines, alimento, energía, "
       "precio de venta, etc.). No borres ni modifiques las celdas verdes o blancas: son fórmulas.")
rr+=1
rr = h(rr, "Paso 3 — Lee tus resultados", 12)
rr = p(rr, "La sección verde «TUS RESULTADOS» se actualiza sola. Fíjate especialmente en:")
rr = p(rr, "• Costo real por kilo → si es mayor que tu precio de venta, PIERDES dinero.", b=True)
rr = p(rr, "• Capital mínimo asegurado → el dinero que debes tener ANTES de sembrar.", b=True)
rr = p(rr, "• Ciclos para recuperar la inversión → cuántas cosechas (~6 meses c/u) hasta recuperar la estructura.", b=True)
rr+=1
rr = h(rr, "Los 3 números que deciden tu rentabilidad", 12)
rr = p(rr, "1. FCA (conversión alimenticia): kg de alimento por kg de pez. Meta ≤ 1,6.")
rr = p(rr, "2. Precio del alimento: es el 40–70% de tu costo. Negocia por volumen.")
rr = p(rr, "3. Precio de venta por kg: consíguelo con el comprador ANTES de sembrar.")
rr+=1
rr = h(rr, "Recuerda", 12, color=YELLOW_D, fillc="FBF3E2")
rr = p(rr, "Los valores precargados son el «Escenario B» del e-book (500 m², 3.000 peces) y son ilustrativos "
       "en USD. Tu realidad depende de los precios de tu país y de tu manejo. Esta planilla es una herramienta "
       "de planeación, no una garantía de resultados.", it if False else 10)

for rw in range(1, rr+1):
    if ws2.row_dimensions[rw].height is None:
        ws2.row_dimensions[rw].height = 16

# Forzar recálculo completo al abrir en Excel / Google Sheets / LibreOffice
try:
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
except Exception:
    wb.calculation.fullCalcOnLoad = True

wb.save("/home/user/psicultura/producto/entregables/Calculadora-ROI-Tilapia.xlsx")
print("XLSX generado OK")
